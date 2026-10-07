import argparse
import datetime
import os
import io
import json
import csv
from flask import Flask, render_template, request, jsonify, session, Response, redirect
from nlp_engine import calculate_match_scores

import PyPDF2
from docx import Document

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'ai-screener-compact-v5-kdu-2026')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_FILE = os.path.join(BASE_DIR, 'results_history.json')
FEEDBACK_FILE = os.path.join(BASE_DIR, 'feedback.json')

def load_results():
    if os.path.exists(RESULTS_FILE):
        try:
            with open(RESULTS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {
        'results': [],
        'stats': {'highly_suitable': 0, 'suitable': 0, 'low_match': 0, 'avg_score': 0.0, 'top_score': 0.0},
        'labels': [],
        'scores': [],
        'model_version': 'v5'
    }

def save_results(data):
    try:
        with open(RESULTS_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2)
    except Exception as e:
        print(f"[v5.0_compact] Failed to save results: {e}")

last_results = load_results()

# ==============================================================================
# TEXT EXTRACTION HELPERS
# ==============================================================================

def extract_text_from_pdf(file_stream):
    """Extract text content from a PDF file stream."""
    file_stream.seek(0, io.SEEK_END)
    size = file_stream.tell()
    file_stream.seek(0)
    if size == 0:
        raise ValueError("Uploaded PDF is empty (0 bytes).")

    reader = PyPDF2.PdfReader(file_stream)
    text = []
    for page in reader.pages:
        try:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
        except Exception:
            continue
    return '\n'.join(text)

def extract_text_from_docx(file_stream):
    """Extract text content from a DOCX file stream."""
    file_stream.seek(0, io.SEEK_END)
    size = file_stream.tell()
    file_stream.seek(0)
    if size == 0:
        raise ValueError("Uploaded DOCX is empty (0 bytes).")

    doc = Document(file_stream)
    text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)
    return '\n'.join(text)

# ==============================================================================
# BENCHMARK TEST SAMPLES (PRESET JOBS & RESUMES)
# ==============================================================================

BENCHMARK_PRESETS = {
    "ai_engineer": {
        "title": "Senior Python & AI Engineer",
        "job_desc": "Looking for a Senior Python & AI Engineer with at least 3 years of experience. Must be proficient in Python, PyTorch or TensorFlow, Natural Language Processing, building REST APIs with FastAPI or Flask, and containerization using Docker and Linux.",
        "candidates": [
            {
                "name": "Alex Chen",
                "role": "Senior AI & Machine Learning Engineer",
                "text": "Senior AI & Machine Learning Engineer with 4 years of experience (2020 - Present). Proficient in Python, PyTorch, TensorFlow, Natural Language Processing (NLP), transformers, Hugging Face, Flask, FastAPI, Docker, and Linux. Built and deployed scalable LLM applications."
            },
            {
                "name": "Sarah Miller",
                "role": "Data Analyst",
                "text": "Data Analyst with 3 years experience (2021 - Present). Skilled in Python, SQL, Pandas, Tableau, and basic Machine Learning algorithms like Scikit-Learn regression and classification. Experienced in data visualization and reporting."
            },
            {
                "name": "David Clark",
                "role": "Junior Web Developer",
                "text": "Junior Web Developer with 1 year experience (2023 - 2024) in HTML, CSS, JavaScript, and basic Python scripting. Passionate about AI and eager to learn machine learning frameworks."
            },
            {
                "name": "Marcus Vance",
                "role": "Executive Chef",
                "text": "Head Chef with 8 years of culinary and kitchen management experience (2016 - 2024). Expert in menu development, food safety protocols, inventory control, and staff supervision."
            }
        ]
    },
    "frontend_react": {
        "title": "Frontend React Developer",
        "job_desc": "Hiring a Frontend Developer with 2+ years experience building responsive web apps using React, JavaScript, TypeScript, HTML5, CSS3, and Redux. Experience with Git and REST APIs required.",
        "candidates": [
            {
                "name": "Elena Rostova",
                "role": "React Developer",
                "text": "Frontend Developer with 3 years of experience (2021 - Present) specializing in React, TypeScript, JavaScript, Redux, HTML5, modern CSS, Git, and integrating RESTful APIs."
            },
            {
                "name": "Kevin Zhang",
                "role": "Full-Stack Node/Vue Engineer",
                "text": "Full-stack engineer with 2 years experience in Vue, JavaScript, Node.js, Express, HTML, CSS, SQL, and Git. Some exposure to React fundamentals."
            },
            {
                "name": "Arthur Pendelton",
                "role": "Civil Construction Engineer",
                "text": "Civil Engineer with 6 years experience managing urban infrastructure builds, concrete testing, safety audits, and CAD site layouts."
            }
        ]
    },
    "devops_cloud": {
        "title": "DevOps & Cloud Infrastructure Engineer",
        "job_desc": "Seeking a Cloud DevOps Engineer with 3+ years experience managing AWS or Azure cloud infrastructure, Docker, Kubernetes, Terraform, CI/CD pipelines with GitHub Actions or Jenkins, and Linux systems administration.",
        "candidates": [
            {
                "name": "Siddharth Rao",
                "role": "Senior Cloud DevOps Engineer",
                "text": "Cloud Infrastructure Engineer with 4 years experience (2020 - Present) in AWS, Docker, Kubernetes, Terraform, Jenkins, GitHub Actions, Linux administration, and microservices telemetry."
            },
            {
                "name": "Liam O'Connor",
                "role": "Linux System Administrator",
                "text": "System Administrator with 3 years experience managing Ubuntu and CentOS Linux servers, Bash scripting, Docker containers, and networking."
            },
            {
                "name": "Hannah Abbott",
                "role": "Digital Marketing Specialist",
                "text": "Marketing lead with 5 years experience managing SEO campaigns, Google Ads, content strategy, and social media brand positioning."
            }
        ]
    }
}

# ==============================================================================
# ROUTES
# ==============================================================================

@app.route('/')
def dashboard():
    """Compact UI single-page dashboard cockpit."""
    global last_results
    last_results = load_results()
    return render_template('dashboard.html',
                           active_page='dashboard',
                           results=last_results['results'],
                           stats=last_results['stats'],
                           labels=last_results['labels'],
                           scores=last_results['scores'],
                           model_version=last_results.get('model_version', 'v5'))

@app.route('/screen')
def screen():
    """Scroll redirect to screening section."""
    return redirect('/#screen')

@app.route('/results')
def results():
    """Scroll redirect to results section."""
    return redirect('/#results')

@app.route('/architecture')
def architecture():
    """Scroll redirect to architecture section."""
    return redirect('/#architecture')

@app.route('/status')
def status():
    """Telemetry health check endpoint."""
    return jsonify({
        "status": "healthy",
        "engine": "v5.0_compact",
        "sbert_weights": "Fine-Tuned Domain SBERT (MNRL)",
        "tfidf_features": 48393,
        "skill_clusters": 4,
        "evaluated_count": len(last_results.get('results', []))
    })

@app.route('/api/presets')
def api_presets():
    """Returns sample benchmark jobs and candidate dossiers for 1-click loading."""
    return jsonify(BENCHMARK_PRESETS)

@app.route('/upload', methods=['POST'])
def upload():
    """Extract clean text from PDF or DOCX file upload."""
    if 'file' not in request.files:
        return jsonify({'status': 'error', 'error': 'No file provided in request'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'status': 'error', 'error': 'No file selected for ingestion'}), 400

    filename = file.filename.lower()
    try:
        if filename.endswith('.pdf'):
            text = extract_text_from_pdf(file.stream)
        elif filename.endswith('.docx'):
            text = extract_text_from_docx(file.stream)
        else:
            return jsonify({'status': 'error', 'error': f'Unsupported format: {filename}. Please use PDF or DOCX.'}), 400

        if not text or not text.strip():
            return jsonify({'status': 'error', 'error': 'No text extracted. File may be rasterized or empty.'}), 400

        return jsonify({'status': 'success', 'text': text, 'filename': file.filename})

    except Exception as e:
        err_msg = str(e)
        if 'empty' in err_msg.lower() or isinstance(e, ValueError):
            return jsonify({'status': 'error', 'error': 'Uploaded file is empty (0 bytes) or corrupted.'}), 400
        return jsonify({'status': 'error', 'error': f'Parsing failure: {err_msg}'}), 500

@app.route('/score', methods=['POST'])
def score():
    """Run AI screening and return candidate ranking as HTMX partial or JSON."""
    global last_results

    data = request.get_json(force=True) if request.is_json else request.form.to_dict()
    job_desc = data.get('job_desc', '')
    resumes = data.get('resumes', [])
    model_version = data.get('model_version', 'v5')

    if isinstance(resumes, str):
        try:
            resumes = json.loads(resumes)
        except Exception:
            resumes = [resumes]

    is_htmx = bool(request.headers.get('HX-Request'))

    def error_response(msg, code=400):
        if is_htmx:
            return f'<div class="p-4 rounded-xl border border-rose-500/30 bg-rose-950/40 text-rose-300 text-sm font-mono">{msg}</div>', code
        return jsonify({'status': 'error', 'error': msg}), code

    if not job_desc or len(job_desc.strip()) < 20:
        return error_response('Job description must be at least 20 characters long.')

    valid_resumes = []
    for r in resumes:
        if isinstance(r, dict):
            txt = r.get('text') or r.get('content') or ''
            if txt.strip():
                valid_resumes.append(r)
        elif isinstance(r, str):
            if r.strip():
                valid_resumes.append(r)
        elif r is not None:
            valid_resumes.append(r)

    if not valid_resumes:
        return error_response('Please provide at least one valid candidate resume with non-empty text.')

    if len(valid_resumes) > 50:
        return error_response('Batch exceeds limit of 50 candidates per execution.')

    # Execute AI Matching Engine
    results_list = calculate_match_scores(job_desc, valid_resumes, model_version=model_version)
    results_list.sort(key=lambda x: x['score'], reverse=True)

    # Calculate Summary Stats
    total_count = len(results_list)
    scores = [r['score'] for r in results_list]
    labels = [r['candidate_id'] for r in results_list]
    avg_score = round(sum(scores) / total_count, 1) if total_count > 0 else 0.0
    top_score = scores[0] if total_count > 0 else 0.0

    stats = {
        'highly_suitable': sum(1 for r in results_list if r['label'] == 'Highly Suitable'),
        'suitable': sum(1 for r in results_list if r['label'] == 'Suitable'),
        'low_match': sum(1 for r in results_list if r['label'] == 'Low Match'),
        'avg_score': avg_score,
        'top_score': top_score
    }

    last_results = {
        'results': results_list,
        'stats': stats,
        'labels': labels,
        'scores': scores,
        'model_version': model_version
    }
    save_results(last_results)

    if request.headers.get('HX-Request'):
        return render_template('partials/results_table.html',
                               results=results_list,
                               stats=stats,
                               labels=labels,
                               scores=scores,
                               model_version=model_version)
    else:
        return jsonify({
            'status': 'success',
            'results': results_list,
            'stats': stats,
            'labels': labels,
            'scores': scores,
            'model_version': model_version
        })

@app.route('/feedback', methods=['POST'])
def feedback():
    """Capture human-in-the-loop recruiter feedback for model calibration."""
    data = request.get_json(force=True)
    candidate_id = data.get('candidate_id')
    feedback_type = data.get('feedback')

    try:
        feedback_history = []
        if os.path.exists(FEEDBACK_FILE):
            with open(FEEDBACK_FILE, 'r', encoding='utf-8') as f:
                feedback_history = json.load(f)

        feedback_history.append({
            'candidate_id': candidate_id,
            'feedback': feedback_type,
            'timestamp': datetime.datetime.now().isoformat()
        })

        with open(FEEDBACK_FILE, 'w', encoding='utf-8') as f:
            json.dump(feedback_history, f, indent=2)
    except Exception as e:
        print(f"[v5.0_compact] Feedback save error: {e}")

    return jsonify({"status": "success", "message": "Feedback recorded for active learning loop"})

@app.route('/export-csv')
def export_csv():
    """Export the evaluated candidate batch as standard CSV."""
    global last_results
    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow([
        'Rank', 'Candidate ID', 'Match Score (%)', 'Label',
        'Experience (Years)', 'SBERT Score (%)', 'TF-IDF Score (%)',
        'Skill Affinity (%)', 'Matched Skills', 'Missing Skills'
    ])

    for idx, r in enumerate(last_results.get('results', []), 1):
        cw.writerow([
            idx,
            r.get('candidate_id', ''),
            r.get('score', 0),
            r.get('label', ''),
            r.get('experience_years', 0),
            r.get('sbert_score', 0),
            r.get('tfidf_score', 0),
            r.get('skill_affinity', 0),
            "; ".join(r.get('matched_skills', [])),
            "; ".join(r.get('missing_skills', []))
        ])

    return Response(
        si.getvalue(),
        mimetype="text/csv",
        headers={"Content-disposition": "attachment; filename=candidate_screening_results_compact.csv"}
    )

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Run v5.0 Compact AI Screener")
    parser.add_argument('--port', type=int, default=int(os.environ.get('PORT', 5001)), help='Port to run server on')
    parser.add_argument('--host', type=str, default='0.0.0.0', help='Host to bind')
    args = parser.parse_args()

    print(f"===========================================================")
    print(f" Group 17 - AI Candidate Screener (v5.0 Compact Edition)")
    print(f" Listening on http://127.0.0.1:{args.port}")
    print(f" Note: v4.1 remains available at http://127.0.0.1:5000")
    print(f"===========================================================")
    app.run(host=args.host, port=args.port, debug=True, use_reloader=False)
