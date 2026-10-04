import datetime
import os
import io
import json
import csv
from flask import Flask, render_template, request, jsonify, session, Response
from nlp_engine import calculate_match_scores

# PDF and DOCX text extraction
import PyPDF2
from docx import Document

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'ai-screener-group17-kdu-2026-dev')

RESULTS_FILE = 'results_history.json'
FEEDBACK_FILE = 'feedback.json'

def load_results():
    if os.path.exists(RESULTS_FILE):
        try:
            with open(RESULTS_FILE, 'r') as f:
                return json.load(f)
        except Exception:
            pass
    return {
        'results': [],
        'stats': {'highly_suitable': 0, 'suitable': 0, 'low_match': 0},
        'labels': [],
        'scores': []
    }

def save_results(data):
    try:
        with open(RESULTS_FILE, 'w') as f:
            json.dump(data, f)
    except Exception as e:
        print(f"Failed to save results: {e}")

# Store last results globally, backed by disk
last_results = load_results()

# ==============================================================================
# TEXT EXTRACTION HELPERS
# ==============================================================================

def extract_text_from_pdf(file_stream):
    """Extract text from a PDF file stream."""
    reader = PyPDF2.PdfReader(file_stream)
    text = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text.append(page_text)
    return '\n'.join(text)

def extract_text_from_docx(file_stream):
    """Extract text from a DOCX file stream."""
    doc = Document(file_stream)
    text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)
    # Also extract from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)
    return '\n'.join(text)

# ==============================================================================
# ROUTES
# ==============================================================================

@app.route('/')
def dashboard():
    """Dashboard / Home page with system overview and stats."""
    return render_template('dashboard.html', active_page='dashboard')

@app.route('/screen')
def screen():
    """Screening page with job description input and resume upload."""
    return render_template('screen.html', active_page='screen')

@app.route('/results')
def results():
    """Results page showing the last screening results."""
    # Reload from disk just in case
    global last_results
    last_results = load_results()
    
    return render_template('results.html',
                           active_page='results',
                           results=last_results['results'],
                           stats=last_results['stats'],
                           labels=last_results['labels'],
                           scores=last_results['scores'])

@app.route('/upload', methods=['POST'])
def upload():
    """Handle file upload — extract text from PDF or DOCX and return it."""
    if 'file' not in request.files:
        return jsonify({'status': 'error', 'error': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'status': 'error', 'error': 'No file selected'}), 400

    filename = file.filename.lower()
    try:
        if filename.endswith('.pdf'):
            text = extract_text_from_pdf(file.stream)
        elif filename.endswith('.docx'):
            text = extract_text_from_docx(file.stream)
        else:
            return jsonify({'status': 'error', 'error': f'Unsupported file type: {filename}'}), 400

        if not text or not text.strip():
            return jsonify({'status': 'error', 'error': 'No text could be extracted from the file. The file may be image-based or empty.'}), 400

        return jsonify({'status': 'success', 'text': text, 'filename': file.filename})

    except Exception as e:
        return jsonify({'status': 'error', 'error': f'Failed to process file: {str(e)}'}), 500

@app.route('/score', methods=['POST'])
def score():
    """Run AI screening and return results as an HTMX partial or JSON."""
    global last_results

    data = request.get_json(force=True)
    job_desc = data.get('job_desc', '')
    resumes = data.get('resumes', [])

    if not job_desc or len(job_desc) < 20:
        return '<div class="text-center py-8 text-red-500 font-medium">Please provide a valid job description (at least 20 characters).</div>', 400
        
    if not resumes:
        return '<div class="text-center py-8 text-red-500 font-medium">Please provide at least one resume.</div>', 400

    if len(resumes) > 50:
        return '<div class="text-center py-8 text-red-500 font-medium">Maximum 50 resumes allowed per batch to prevent server overload.</div>', 400

    # Run the AI matching engine
    results = calculate_match_scores(job_desc, resumes)

    # Sort by score descending
    results.sort(key=lambda x: x['score'], reverse=True)

    # Calculate stats
    stats = {
        'highly_suitable': sum(1 for r in results if r['label'] == 'Highly Suitable'),
        'suitable': sum(1 for r in results if r['label'] == 'Suitable'),
        'low_match': sum(1 for r in results if r['label'] == 'Low Match'),
    }

    labels = [r['candidate_id'] for r in results]
    scores = [r['score'] for r in results]

    # Store globally and on disk for the results page
    last_results = {
        'results': results,
        'stats': stats,
        'labels': labels,
        'scores': scores
    }
    save_results(last_results)

    # Check if HTMX request (returns HTML partial) or regular API call (returns JSON)
    if request.headers.get('HX-Request'):
        return render_template('partials/results_table.html',
                               results=results,
                               stats=stats,
                               labels=labels,
                               scores=scores)
    else:
        return jsonify(results)

@app.route('/feedback', methods=['POST'])
def feedback():
    """Capture human-in-the-loop feedback for future model retraining."""
    data = request.get_json(force=True)
    candidate_id = data.get('candidate_id')
    feedback_type = data.get('feedback')

    print(f"Human-in-the-loop Feedback Received: {candidate_id} -> {feedback_type}")
    
    # Save to disk for ML retraining pipeline
    try:
        feedback_history = []
        if os.path.exists(FEEDBACK_FILE):
            with open(FEEDBACK_FILE, 'r') as f:
                feedback_history = json.load(f)
        
        feedback_history.append({
            'candidate_id': candidate_id,
            'feedback': feedback_type,
            'timestamp': str(datetime.datetime.now()) if 'datetime' in globals() else 'now' # quick fix for timestamp
        })
        
        with open(FEEDBACK_FILE, 'w') as f:
            json.dump(feedback_history, f)
    except Exception as e:
        print(f"Error saving feedback: {e}")

    return jsonify({"status": "success", "message": "Feedback recorded for future ML training!"})

@app.route('/export-csv')
def export_csv():
    """Download the latest results as a CSV file."""
    global last_results
    
    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow(['Candidate ID', 'Score (%)', 'Label', 'SBERT Score', 'TF-IDF Score', 'Skill Affinity', 'Matched Skills', 'Missing Skills', 'Experience (Years)'])
    
    for r in last_results['results']:
        cw.writerow([
            r.get('candidate_id', ''),
            r.get('score', 0),
            r.get('label', ''),
            r.get('sbert_score', 0),
            r.get('tfidf_score', 0),
            r.get('skill_affinity', 0),
            ", ".join(r.get('matched_skills', [])),
            ", ".join(r.get('missing_skills', [])),
            r.get('experience_years', 0)
        ])
        
    return Response(
        si.getvalue(),
        mimetype="text/csv",
        headers={"Content-disposition": "attachment; filename=screening_results.csv"}
    )

# ==============================================================================
# LEGACY ROUTE (keep old index.html working)
# ==============================================================================

@app.route('/legacy')
def legacy():
    """Original V1 prototype interface (preserved for reference)."""
    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True)
